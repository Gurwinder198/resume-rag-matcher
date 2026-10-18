"""Generate a synthetic dataset: 30+ diverse resumes and 5+ job descriptions.

Run:  python scripts/generate_dataset.py [--out-dir .]
Needs: reportlab, python-docx (both in requirements.txt)

Writes:
    resumes/*.txt|.pdf|.docx
    job_descriptions/*.txt
    dataset_manifest.json  (ground truth: file -> role, for evaluation)
"""

from __future__ import annotations

import argparse
import json
from datetime import datetime
from pathlib import Path

CURRENT_YEAR = datetime.now().year

ROLES = {
    "ML Engineer": {
        "skills": ["Python", "PyTorch", "TensorFlow", "scikit-learn", "Machine Learning", "NLP", "Docker", "AWS"],
        "title": "Machine Learning Engineer",
    },
    "Data Engineer": {
        "skills": ["Python", "Spark", "Airflow", "dbt", "SQL", "Snowflake", "Kafka"],
        "title": "Data Engineer",
    },
    "Data Scientist": {
        "skills": ["Python", "Pandas", "NumPy", "SQL", "Machine Learning", "Tableau"],
        "title": "Data Scientist",
    },
    "Frontend Developer": {
        "skills": ["JavaScript", "TypeScript", "React", "Next.js", "Jest", "Figma"],
        "title": "Frontend Developer",
    },
    "Backend Developer": {
        "skills": ["Java", "Spring", "SQL", "Docker", "Kubernetes", "REST", "Kafka"],
        "title": "Backend Developer",
    },
    "DevOps Engineer": {
        "skills": ["Kubernetes", "Terraform", "AWS", "Bash", "CI/CD", "Linux", "Docker"],
        "title": "DevOps Engineer",
    },
    "Full Stack Developer": {
        "skills": ["JavaScript", "React", "Node.js", "Python", "Django", "SQL"],
        "title": "Full Stack Developer",
    },
    "QA Engineer": {
        "skills": ["Selenium", "Python", "Jest", "CI/CD", "Jira", "Git"],
        "title": "QA Automation Engineer",
    },
}

FIRST_NAMES = [
    "Alice", "Bob", "Carla", "David", "Elena", "Farhan", "Grace", "Hiro",
    "Isabel", "Jamal", "Kavya", "Liam", "Maya", "Noah", "Olivia", "Priya",
    "Quentin", "Rosa", "Sam", "Tara", "Umar", "Vera", "Wei", "Ximena",
    "Yusuf", "Zoe", "Anders", "Bianca", "Chen", "Dalia", "Ewan", "Fiona",
]
LAST_NAMES = [
    "Nguyen", "Smith", "Mehta", "Okafor", "Rossi", "Khan", "Lindqvist", "Tanaka",
    "Garcia", "Haddad", "Rao", "Murphy", "Singh", "Clark", "Dubois", "Sharma",
    "Reyes", "Fernandez", "Lee", "Patel", "Hassan", "Ivanova", "Zhang", "Torres",
    "Ahmed", "Brooks", "Larsen", "Costa", "Wu", "Hernandez", "Fraser", "Moreau",
]

COMPANIES = ["Acme", "Globex", "Initech", "Umbrella", "Soylent", "Hooli",
             "Vandelay", "Stark Industries", "Wayne Enterprises", "Wonka Inc"]
SCHOOLS = ["Stanford", "MIT", "IIT Delhi", "UCLA", "University of Toronto",
           "NUS", "ETH Zurich", "Georgia Tech"]
DEGREES = ["BSc Computer Science", "MSc Computer Science", "BTech Information Technology",
           "BEng Software Engineering", "MSc Data Science"]

YEARS_CYCLE = [2, 3, 5, 7]


def _resume_text(name: str, role_title: str, skills: list[str], years: int,
                  company: str, school: str, degree: str) -> str:
    return f"""{name}
{role_title} | {name.lower().replace(' ', '.')}@example.com

SUMMARY
{role_title} with {years} years of experience in {skills[0]} and {skills[1]}.

SKILLS
{', '.join(skills)}

EXPERIENCE
{role_title}, {company} ({CURRENT_YEAR - years} - present)
- Delivered production {skills[0]} systems used across the organization.
- Collaborated cross-functionally to ship features on schedule.

EDUCATION
{degree}, {school} ({CURRENT_YEAR - years - 4})
"""


def _write_pdf(path: Path, text: str) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(path), pagesize=A4)
    y = 800
    for line in text.splitlines():
        is_heading = bool(line.strip()) and line.isupper()
        c.setFont("Helvetica-Bold" if is_heading else "Helvetica", 11)
        c.drawString(60, y, line[:95])
        y -= 16
        if y < 40:
            c.showPage()
            y = 800
    c.save()


def _write_docx(path: Path, text: str) -> None:
    import docx

    d = docx.Document()
    for line in text.splitlines():
        if line.strip() and line.isupper():
            d.add_paragraph(line, style="Heading 1")
        else:
            d.add_paragraph(line)
    d.save(str(path))


def generate_resumes(out_dir: Path) -> list[dict]:
    out_dir.mkdir(parents=True, exist_ok=True)
    manifest = []
    index = 0
    for role, info in ROLES.items():
        for person in range(4):
            first = FIRST_NAMES[index % len(FIRST_NAMES)]
            last = LAST_NAMES[index % len(LAST_NAMES)]
            name = f"{first} {last}"
            years = YEARS_CYCLE[person % len(YEARS_CYCLE)]
            skills = info["skills"][: 5 + (person % 3)]
            company = COMPANIES[index % len(COMPANIES)]
            school = SCHOOLS[index % len(SCHOOLS)]
            degree = DEGREES[index % len(DEGREES)]
            text = _resume_text(name, info["title"], skills, years, company, school, degree)

            slug = f"resume_{first.lower()}_{last.lower()}"
            fmt = index % 8
            if fmt == 0:
                filename = f"{slug}.pdf"
                _write_pdf(out_dir / filename, text)
            elif fmt == 1:
                filename = f"{slug}.docx"
                _write_docx(out_dir / filename, text)
            else:
                filename = f"{slug}.txt"
                (out_dir / filename).write_text(text, encoding="utf-8")

            manifest.append({
                "file": filename, "name": name, "role": role,
                "years_experience": years, "skills": skills,
            })
            index += 1
    return manifest


JD_TEMPLATE = """We are hiring a {title}.

About the role:
We need someone who can own {skill_a} and {skill_b} work end to end.

Must have:
{skills}

Requires {years}+ years of experience.
"""

JD_ROLES = ["ML Engineer", "Data Engineer", "Frontend Developer", "DevOps Engineer", "Full Stack Developer"]


def generate_job_descriptions(out_dir: Path) -> list[dict]:
    out_dir.mkdir(parents=True, exist_ok=True)
    manifest = []
    for role in JD_ROLES:
        info = ROLES[role]
        skills = info["skills"][:4]
        text = JD_TEMPLATE.format(
            title=info["title"], skill_a=skills[0], skill_b=skills[1],
            skills=", ".join(skills), years=5,
        )
        filename = f"jd_{role.lower().replace(' ', '_')}.txt"
        (out_dir / filename).write_text(text, encoding="utf-8")
        manifest.append({"file": filename, "role": role, "min_years": 5, "required_skills": skills})
    return manifest


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Generate the synthetic resume/JD dataset.")
    parser.add_argument("--out-dir", default=".")
    args = parser.parse_args(argv)

    base = Path(args.out_dir)
    resumes = generate_resumes(base / "resumes")
    jds = generate_job_descriptions(base / "job_descriptions")
    manifest = {"resumes": resumes, "job_descriptions": jds}
    (base / "dataset_manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")

    print(f"Wrote {len(resumes)} resume(s) and {len(jds)} job description(s) to {base}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
